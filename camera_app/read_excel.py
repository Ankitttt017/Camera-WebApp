import openpyxl
import json
from pathlib import Path

file_path = Path(r'c:\Vikram Dubey One Drive\OneDrive - ricoauto.in\Desktop\Project\AB\Camera-WebApp\Downtime_Hierarchy_Final.xlsx')

wb = openpyxl.load_workbook(file_path, data_only=True)
print("Sheet names:", wb.sheetnames)

hierarchy = {}

for sheet_name in wb.sheetnames:
    if sheet_name.lower() == 'overview':
        continue
        
    sheet = wb[sheet_name]
    l1_name = sheet_name.strip()
    hierarchy[l1_name] = {}
    
    # Assuming columns: Sr No, Category, Sub-Category
    # Let's find the headers first
    category_col = None
    sub_category_col = None
    
    # Find header row
    header_row = 1
    for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
        row_strs = [str(x).lower().strip() if x else '' for x in row]
        if any('category' in x for x in row_strs):
            for i, val in enumerate(row_strs):
                if 'sub' in val and 'category' in val:
                    sub_category_col = i
                elif 'category' in val:
                    category_col = i
            break
        header_row += 1
        
    if category_col is None or sub_category_col is None:
        print(f"Skipping {sheet_name}: Could not find Category or Sub-Category columns.")
        continue
        
    for row in sheet.iter_rows(min_row=header_row+1, values_only=True):
        cat = str(row[category_col]).strip() if row[category_col] else None
        sub_cat = str(row[sub_category_col]).strip() if row[sub_category_col] else None
        
        if not cat or cat.lower() == 'none' or not sub_cat or sub_cat.lower() == 'none':
            continue
            
        if cat not in hierarchy[l1_name]:
            hierarchy[l1_name][cat] = []
            
        if sub_cat not in hierarchy[l1_name][cat]:
            hierarchy[l1_name][cat].append(sub_cat)

with open(r'c:\Vikram Dubey One Drive\OneDrive - ricoauto.in\Desktop\Project\AB\Camera-WebApp\camera_app\frontend\src\downtimeData.ts', 'w', encoding='utf-8') as f:
    f.write("// This file contains the hierarchy for Downtime reasons.\n")
    f.write("// Level 1 (L1) -> Downtime Type\n")
    f.write("// Level 2 (L2) -> Category\n")
    f.write("// Level 3 (L3) -> Sub-Category\n\n")
    f.write("export const ACTUAL_DOWNTIME_DATA: Record<string, Record<string, string[]>> = ")
    f.write(json.dumps(hierarchy, indent=2))
    f.write(";\n")
    
print("Successfully generated downtimeData.ts!")
